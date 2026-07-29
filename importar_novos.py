#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importa novos funcionários a partir de arquivos soltos na pasta "entrada/"
(.xlsx, .xls, .csv, .docx, .pdf) e mescla na planilha mestre
controle_experiencias.xlsx, sem duplicar quem já está cadastrado.

Regras:
  - Identifica colunas por palavras-chave (nome, setor, contratação/admissão,
    30 dias, 60 dias), então funciona mesmo que o arquivo novo não tenha
    exatamente os mesmos nomes de coluna da planilha mestre.
  - Se o arquivo trouxer só a data de contratação, calcula 30 e 60 dias
    automaticamente (contratação + 30 / + 60 dias corridos).
  - Duplicado = mesmo nome (ignorando maiúsc./minúsc., acento e espaços
    extras) já presente na planilha mestre. Esses são ignorados.
  - Depois de importar, move o arquivo de entrada/ para processados/
    (prefixado com data/hora) para não ser lido de novo.

Uso:
  python3 importar_novos.py [pasta_base]
  (pasta_base default = diretório atual; espera entrada/, processados/ e
   controle_experiencias.xlsx dentro dela)
"""
import sys
import os
import re
import shutil
import unicodedata
from datetime import datetime, date, timedelta

import openpyxl
import pandas as pd

BASE = sys.argv[1] if len(sys.argv) > 1 else "."
MASTER = os.path.join(BASE, "controle_experiencias.xlsx")
ENTRADA = os.path.join(BASE, "entrada")
PROCESSADOS = os.path.join(BASE, "processados")

HEADER_MAP = {
    "nome": "nome",
    "setor": "setor",
    "contrat": "contratacao",
    "admiss": "contratacao",
    "30": "d30",
    "60": "d60",
}


def norm_text(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return s


def norm_name(s):
    s = norm_text(s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def guess_column(header_cell):
    h = norm_text(header_cell)
    if not h:
        return None
    if "nome" in h:
        return "nome"
    if "setor" in h:
        return "setor"
    if "contrat" in h or "admiss" in h:
        return "contratacao"
    if "30" in h and "dia" in h:
        return "d30"
    if "60" in h and "dia" in h:
        return "d60"
    return None


def parse_date_flex(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    # tenta extrair "dd/mm/aaaa" de dentro de um texto maior
    m = re.search(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    if m:
        d, mo, y = m.groups()
        y = int(y)
        if y < 100:
            y += 2000
        try:
            return date(y, int(mo), int(d))
        except ValueError:
            return None
    return None


def table_to_records(rows):
    """rows: lista de listas (linha 0 = cabeçalho). Retorna lista de dicts
    com chaves nome/setor/contratacao/d30/d60 (o que existir)."""
    if not rows:
        return []
    header = rows[0]
    colmap = {}
    for i, cell in enumerate(header):
        key = guess_column(cell)
        if key:
            colmap[i] = key
    if "nome" not in colmap.values():
        return []

    records = []
    for row in rows[1:]:
        rec = {}
        for i, val in enumerate(row):
            key = colmap.get(i)
            if not key:
                continue
            if key in ("contratacao", "d30", "d60"):
                rec[key] = parse_date_flex(val)
            else:
                rec[key] = str(val).strip() if val is not None else ""
        if rec.get("nome"):
            records.append(rec)
    return records


def extract_from_xlsx_csv(path):
    if path.lower().endswith(".csv"):
        df = pd.read_csv(path, header=None, dtype=str)
    else:
        df = pd.read_excel(path, header=None, dtype=str, engine="openpyxl" if path.lower().endswith(("xlsx", "xlsm")) else None)
    rows = df.fillna("").values.tolist()
    # acha a linha de cabeçalho: primeira linha que tenha "nome"
    header_idx = None
    for i, row in enumerate(rows):
        if any(guess_column(c) == "nome" for c in row):
            header_idx = i
            break
    if header_idx is None:
        return []
    return table_to_records(rows[header_idx:])


def extract_from_docx(path):
    from docx import Document
    doc = Document(path)
    records = []
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        records.extend(table_to_records(rows))
    return records


def extract_from_pdf(path):
    import pdfplumber
    records = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                records.extend(table_to_records(table))
    return records


def extract(path):
    low = path.lower()
    try:
        if low.endswith((".xlsx", ".xls", ".xlsm", ".csv")):
            return extract_from_xlsx_csv(path)
        if low.endswith(".docx"):
            return extract_from_docx(path)
        if low.endswith(".pdf"):
            return extract_from_pdf(path)
    except Exception as e:
        print(f"  [erro ao ler {os.path.basename(path)}: {e}]")
        return []
    print(f"  [formato não suportado: {os.path.basename(path)}]")
    return []


def load_master_names(master_path):
    wb = openpyxl.load_workbook(master_path)
    ws = wb.active
    names = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            names.add(norm_name(row[0]))
    return wb, ws, names


def main():
    if not os.path.isdir(ENTRADA):
        os.makedirs(ENTRADA, exist_ok=True)
    if not os.path.isdir(PROCESSADOS):
        os.makedirs(PROCESSADOS, exist_ok=True)
    if not os.path.exists(MASTER):
        print(f"Planilha mestre não encontrada em {MASTER}.")
        return

    files = [f for f in sorted(os.listdir(ENTRADA))
             if not f.startswith(".") and os.path.isfile(os.path.join(ENTRADA, f))]

    if not files:
        print("Nenhum arquivo novo em entrada/.")
        return

    wb, ws, existing_names = load_master_names(MASTER)

    total_new = 0
    total_dup = 0
    total_files_ok = 0

    for fname in files:
        fpath = os.path.join(ENTRADA, fname)
        print(f"Lendo: {fname}")
        records = extract(fpath)
        if not records:
            print(f"  Nenhum registro reconhecido em {fname} (não foi movido; verifique o formato).")
            continue

        added_here = 0
        for rec in records:
            nome = rec.get("nome", "").strip()
            if not nome:
                continue
            key = norm_name(nome)
            if key in existing_names:
                total_dup += 1
                continue

            setor = rec.get("setor", "")
            contratacao = rec.get("contratacao")
            d30 = rec.get("d30") or (contratacao + timedelta(days=30) if contratacao else None)
            d60 = rec.get("d60") or (contratacao + timedelta(days=60) if contratacao else None)

            ws.append([
                nome,
                setor,
                contratacao.strftime("%d/%m/%Y") if contratacao else "",
                d30.strftime("%d/%m/%Y") if d30 else "",
                d60.strftime("%d/%m/%Y") if d60 else "",
            ])
            existing_names.add(key)
            added_here += 1
            total_new += 1

        print(f"  {added_here} novo(s) adicionado(s), "
              f"{len(records) - added_here} já existiam (ignorados).")

        # move para processados/
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = os.path.join(PROCESSADOS, f"{ts}_{fname}")
        shutil.move(fpath, dest)
        total_files_ok += 1

    if total_new > 0:
        wb.save(MASTER)

    print(f"\nResumo: {total_files_ok} arquivo(s) processado(s), "
          f"{total_new} funcionário(s) novo(s) adicionado(s), "
          f"{total_dup} duplicado(s) ignorado(s).")


if __name__ == "__main__":
    main()
